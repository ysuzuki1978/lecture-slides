# /// script
# requires-python = ">=3.10"
# dependencies = ["pillow>=10", "python-pptx>=0.6.23"]
# ///
"""スライド（PDF または PPTX）を docs/decks/<id>/ に変換し、docs/decks.json に登録する。

  uv run tools/publish_deck.py deck.pptx --id peep-heart-lung --title "…" --subtitle "…"
  uv run tools/publish_deck.py deck.pdf  --id peep-heart-lung --title "…" --notes-from deck.pptx
  uv run tools/publish_deck.py --remove peep-heart-lung

PPTX は LibreOffice で PDF にする（和文フォントは Noto で代替されるので、
見た目を厳密にしたいときは PowerPoint で書き出した PDF を渡し、ノートは --notes-from で PPTX から取る）。
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
DECKS = DOCS / "decks"
CATALOG = DOCS / "decks.json"
FONTS_CONF = Path(__file__).resolve().parent / "fonts.conf"
DEFAULT_SOFFICE = Path.home() / "scratch/libreoffice/opt/libreoffice25.8/program/soffice"

WIDTH = 1920       # スライド画像の横幅（px）
THUMB = 480        # サムネイルの横幅（px）
QUALITY = 82       # WebP 品質


def die(msg: str) -> None:
    print(f"error: {msg}", file=sys.stderr)
    sys.exit(1)


def soffice_path() -> str:
    for cand in (os.environ.get("SOFFICE"), shutil.which("soffice"), str(DEFAULT_SOFFICE)):
        if cand and Path(cand).exists():
            return cand
    die("LibreOffice（soffice）が見つからない。SOFFICE 環境変数で場所を指定するか、PDF を渡す")
    return ""


def pptx_to_pdf(pptx: Path, out: Path) -> Path:
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    if FONTS_CONF.exists():
        env["FONTCONFIG_FILE"] = str(FONTS_CONF)
    profile = out / "lo_profile"
    cmd = [soffice_path(), "--headless", "--norestore", f"-env:UserInstallation={profile.as_uri()}",
           "--convert-to", "pdf", "--outdir", str(out), str(pptx)]
    pdf = out / (pptx.stem + ".pdf")
    # 新しいプロファイルの初回起動は変換せずに終わることがあるので 1 回だけやり直す
    for _ in range(2):
        subprocess.run(cmd, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
        if pdf.exists():
            return pdf
    die("PPTX から PDF への変換に失敗した")
    return pdf


def page_count(pdf: Path) -> int:
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, check=True).stdout
    m = re.search(r"^Pages:\s+(\d+)", info, re.M)
    if not m:
        die("pdfinfo でページ数を読めない")
    return int(m.group(1))


def page_text(pdf: Path, n: int) -> str:
    txt = subprocess.run(["pdftotext", "-f", str(n), "-l", str(n), "-enc", "UTF-8", str(pdf), "-"],
                         capture_output=True, text=True).stdout
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt[:800]


def pptx_notes(pptx: Path) -> list[str]:
    from pptx import Presentation

    notes = []
    for slide in Presentation(str(pptx)).slides:
        text = ""
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
        notes.append(text)
    return notes


def load_catalog() -> dict:
    if CATALOG.exists():
        return json.loads(CATALOG.read_text(encoding="utf-8"))
    return {"decks": []}


def save_catalog(cat: dict) -> None:
    cat["decks"].sort(key=lambda d: d.get("date", ""), reverse=True)
    CATALOG.write_text(json.dumps(cat, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def remove(deck_id: str) -> None:
    cat = load_catalog()
    before = len(cat["decks"])
    cat["decks"] = [d for d in cat["decks"] if d["id"] != deck_id]
    shutil.rmtree(DECKS / deck_id, ignore_errors=True)
    save_catalog(cat)
    print(f"removed {deck_id} ({before - len(cat['decks'])} entry)")


def publish(a: argparse.Namespace) -> None:
    src = Path(a.source).resolve()
    if not src.exists():
        die(f"{src} がない")
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", a.id):
        die("--id は半角英小文字・数字・ハイフンのみ（例：peep-heart-lung）")

    with tempfile.TemporaryDirectory(prefix="publish_") as tmpd:
        tmp = Path(tmpd)
        notes_src = Path(a.notes_from).resolve() if a.notes_from else None
        if src.suffix.lower() == ".pptx":
            pdf = pptx_to_pdf(src, tmp)
            notes_src = notes_src or src
        elif src.suffix.lower() == ".pdf":
            pdf = src
        else:
            die("PDF か PPTX を渡す")

        n = page_count(pdf)
        notes = pptx_notes(notes_src) if notes_src else []
        if notes and len(notes) != n:
            print(f"warning: PDF は {n} ページ、PPTX は {len(notes)} 枚。ノートは先頭から順に対応させる"
                  "（非表示スライドがあるとずれる）", file=sys.stderr)

        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(WIDTH), "-scale-to-y", "-1",
                        str(pdf), str(tmp / "p")], check=True)
        pngs = sorted(tmp.glob("p-*.png"))
        if len(pngs) != n:
            die(f"画像化したページ数（{len(pngs)}）が PDF（{n}）と合わない")

        out = DECKS / a.id
        shutil.rmtree(out, ignore_errors=True)
        out.mkdir(parents=True)

        slides = []
        aspect = 16 / 9
        for i, png in enumerate(pngs, start=1):
            im = Image.open(png).convert("RGB")
            aspect = im.width / im.height
            name, thumb = f"s{i:02d}.webp", f"t{i:02d}.webp"
            im.save(out / name, "WEBP", quality=QUALITY, method=6)
            im.resize((THUMB, round(THUMB / aspect)), Image.LANCZOS).save(out / thumb, "WEBP", quality=78, method=6)
            slides.append({
                "img": name,
                "thumb": thumb,
                "text": page_text(pdf, i),
                "notes": notes[i - 1] if i - 1 < len(notes) else "",
            })

        rev = hashlib.sha256(pdf.read_bytes()).hexdigest()[:10]
        date = a.date or dt.date.today().isoformat()
        deck = {"id": a.id, "title": a.title, "subtitle": a.subtitle or "", "date": date,
                "rev": rev, "aspect": round(aspect, 4), "slides": slides}
        (out / "deck.json").write_text(json.dumps(deck, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    cat = load_catalog()
    cat["decks"] = [d for d in cat["decks"] if d["id"] != a.id]
    cat["decks"].append({"id": a.id, "title": a.title, "subtitle": a.subtitle or "", "date": date,
                         "slides": len(slides), "thumb": "t01.webp", "aspect": round(aspect, 4), "rev": rev})
    save_catalog(cat)

    size = sum(f.stat().st_size for f in out.iterdir()) / 1e6
    with_notes = sum(1 for s in slides if s["notes"])
    print(f"published {a.id}: {len(slides)} slides, notes {with_notes}, {size:.1f} MB, rev {rev}")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("source", nargs="?", help="PDF または PPTX")
    p.add_argument("--id", help="URL に使う名前（例：peep-heart-lung）")
    p.add_argument("--title", help="一覧に出すタイトル")
    p.add_argument("--subtitle", help="一覧に出す副題")
    p.add_argument("--date", help="YYYY-MM-DD（省略時は今日）")
    p.add_argument("--notes-from", help="PDF を渡すとき、発表者ノートを取り出す PPTX")
    p.add_argument("--remove", metavar="ID", help="登録を削除する")
    a = p.parse_args()
    if a.remove:
        remove(a.remove)
        return
    if not (a.source and a.id and a.title):
        p.error("source, --id, --title は必須")
    publish(a)


if __name__ == "__main__":
    main()
